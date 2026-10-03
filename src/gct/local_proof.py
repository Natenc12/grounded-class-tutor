"""Bounded JSON-lines peer caller for the local desktop proof.

Run ``python -m gct.local_proof [--inspect]``. The caller supplies one selected
file (or an explicitly synthetic sample), then answers ``generate`` messages.
No database, embedder, provider client, or semantic retrieval runs here. The
existing parser, chunker, and Grounder retain ownership of provenance and answer
validation. Selected chunks receive a constant score, not a relevance claim.
"""

from __future__ import annotations

import argparse
import contextlib
import hashlib
import json
import logging
import os
import stat
import sys
import zipfile
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import TYPE_CHECKING, TextIO

if TYPE_CHECKING:
    from collections.abc import Sequence

    from gct.ingest.parse import ParsedUnit
    from gct.providers.base import Message
    from gct.retriever.retrieve import RetrievedChunk

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_PAGES = 500
MAX_SELECTED_PAGES = 5
MAX_TEXT_CHARS = 40_000
MAX_QUESTION_CHARS = 2_000
MAX_LINE_CHARS = 262_144
# Compressed file size does not bound the parser's allocations. These limits are
# specific to the desktop proof, not a change to the production ingest contract.
MAX_PPTX_MEMBERS = 5_000
MAX_PPTX_MEMBER_BYTES = 32 * 1024 * 1024
MAX_PPTX_EXPANDED_BYTES = 64 * 1024 * 1024
MAX_PDF_STREAM_BYTES = 16 * 1024 * 1024
MAX_PDF_DECODED_BYTES = 64 * 1024 * 1024
_UNSET = object()
SAMPLE_FILENAME = "local-proof-sample.pdf"
SAMPLE_QUESTION = "What is active recall, and how does spaced practice work?"
SAMPLE_TEXT = (
    "Synthetic educational sample created for the local Grounded Class Tutor proof. "
    "Active recall means retrieving information from memory instead of only rereading it. "
    "Spaced practice spreads study sessions over time instead of putting them all in one session. "
    "This sample does not specify an ideal interval between study sessions."
)


class ProofError(Exception):
    """A closed, safe error message; never constructed from parser/provider text."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


class GenerationFailure(Exception):
    """Terminal to the Grounder: transport/SDK errors do not consume a retry."""


class _DiscardOutput:
    """Suppress third-party parser diagnostics that could echo document contents."""

    def write(self, text: str) -> int:
        return len(text)

    def flush(self) -> None:
        pass


@dataclass(frozen=True)
class Document:
    filename: str
    page_count: int
    units: list[ParsedUnit]
    synthetic: bool = False


def _read_object(stream: TextIO) -> dict:
    line = stream.readline(MAX_LINE_CHARS + 1)
    if not line or len(line) > MAX_LINE_CHARS:
        raise ProofError("invalid_protocol", "Expected a bounded JSON object on one line.")
    try:
        value = json.loads(line)
    except (ValueError, RecursionError) as exc:
        raise ProofError("invalid_protocol", "Could not read the JSON request.") from exc
    if not isinstance(value, dict):
        raise ProofError("invalid_protocol", "The request must be a JSON object.")
    return value


def _emit(stream: TextIO, value: dict) -> None:
    stream.write(json.dumps(value, ensure_ascii=True, allow_nan=False) + "\n")
    stream.flush()


def _check_pptx_archive(path: Path) -> None:
    """Reject oversized ZIP packages before python-pptx eagerly loads their parts."""
    with path.open("rb") as stream:
        if stream.read(8) == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
            raise ProofError("protected", "Password-protected files are not supported.")
    with zipfile.ZipFile(path) as archive:
        members = archive.infolist()
        if len(members) > MAX_PPTX_MEMBERS:
            raise ProofError("too_long", "The presentation contains too many archive parts.")
        if any(member.flag_bits & 1 for member in members):
            raise ProofError("protected", "Password-protected files are not supported.")
        if len({member.filename for member in members}) != len(members):
            raise ProofError("unparseable", "The presentation contains duplicate archive parts.")
        if (
            any(member.file_size > MAX_PPTX_MEMBER_BYTES for member in members)
            or sum(member.file_size for member in members) > MAX_PPTX_EXPANDED_BYTES
        ):
            raise ProofError("too_long", "The presentation expands beyond this proof's limits.")


@contextlib.contextmanager
def _pdf_decoding_limits():
    """Apply a decode budget inside the dedicated, single-request parser process.

    pypdf's bounded decoders stop an individual stream before unbounded expansion;
    counting decoded streams also bounds the cumulative decoded data. This is not
    an OS memory sandbox: parser object/operation overhead still needs the caller's
    timeout and process cleanup. Restore globals for in-process test callers.
    """
    from pypdf import filters
    from pypdf.errors import LimitReachedError

    names = (
        "MAX_DECLARED_STREAM_LENGTH",
        "MAX_ARRAY_BASED_STREAM_OUTPUT_LENGTH",
        "JBIG2_MAX_OUTPUT_LENGTH",
        "LZW_MAX_OUTPUT_LENGTH",
        "RUN_LENGTH_MAX_OUTPUT_LENGTH",
        "ZLIB_MAX_OUTPUT_LENGTH",
        "FLATE_MAX_BUFFER_SIZE",
    )
    original_limits = {name: getattr(filters, name) for name in names}
    original_decode = filters.decode_stream_data
    decoded_bytes = 0

    def decode(stream):
        nonlocal decoded_bytes
        remaining = MAX_PDF_DECODED_BYTES - decoded_bytes
        if remaining <= 0:
            raise LimitReachedError("The proof's PDF decode budget was exceeded.")
        for name, original in original_limits.items():
            setattr(filters, name, min(original, MAX_PDF_STREAM_BYTES, remaining))
        data = original_decode(stream)
        decoded_bytes += len(data)
        if len(data) > MAX_PDF_STREAM_BYTES or decoded_bytes > MAX_PDF_DECODED_BYTES:
            raise LimitReachedError("The proof's PDF decode budget was exceeded.")
        return data

    try:
        for name, original in original_limits.items():
            setattr(filters, name, min(original, MAX_PDF_STREAM_BYTES))
        filters.decode_stream_data = decode
        yield
    finally:
        filters.decode_stream_data = original_decode
        for name, original in original_limits.items():
            setattr(filters, name, original)


def _load_document(request: dict) -> Document:
    from pptx import Presentation
    from pypdf import PdfReader
    from pypdf.errors import LimitReachedError

    from gct.ingest.parse import ParsedUnit, ParseError, parse_file

    if request.get("sample") is True and "file_path" not in request:
        return Document(
            SAMPLE_FILENAME, 1, [ParsedUnit(SAMPLE_TEXT, SAMPLE_FILENAME, 1)], synthetic=True
        )
    value = request.get("file_path")
    if "sample" in request or not isinstance(value, str) or not value or "\x00" in value:
        raise ProofError("invalid_file", "Choose one PDF or PPTX file, or the synthetic sample.")
    path = Path(value)
    if path.suffix.lower() not in {".pdf", ".pptx"}:
        raise ProofError("unsupported", "Only PDF and PPTX files are supported.")
    try:
        info = path.stat()
        if not stat.S_ISREG(info.st_mode):
            raise ProofError("invalid_file", "Choose a regular PDF or PPTX file.")
        if info.st_size > MAX_FILE_BYTES:
            raise ProofError("file_too_large", "The selected file exceeds the 10 MiB limit.")
        with contextlib.ExitStack() as limits:
            # Count physical pages first: parse_file deliberately omits empty pages,
            # which must not renumber either the picker or the citation spine.
            if path.suffix.lower() == ".pdf":
                limits.enter_context(_pdf_decoding_limits())
                reader = PdfReader(path)
                if reader.is_encrypted:
                    raise ProofError("protected", "Password-protected files are not supported.")
                page_count = len(reader.pages)
            else:
                _check_pptx_archive(path)
                page_count = len(Presentation(str(path)).slides)
            if page_count > MAX_PAGES:
                raise ProofError(
                    "too_many_pages", "This proof accepts at most 500 pages or slides."
                )
            try:
                units = parse_file(path)
            except ParseError as exc:
                if exc.reason == "empty":
                    units = []
                else:
                    raise
    except ProofError:
        raise
    except LimitReachedError as exc:
        raise ProofError("too_long", "The PDF expands beyond this proof's limits.") from exc
    except ParseError as exc:
        messages = {
            "protected": "Password-protected files are not supported.",
            "unsupported": "Only PDF and PPTX files are supported.",
            "unparseable": "The selected document could not be read.",
            "too_long": "The selected document exceeds this proof's parser limits.",
        }
        raise ProofError(
            exc.reason, messages.get(exc.reason, "The document could not be read.")
        ) from exc
    except OSError as exc:
        raise ProofError("file_unavailable", "The selected file could not be opened.") from exc
    except Exception as exc:
        raise ProofError("unparseable", "The selected document could not be read.") from exc
    return Document(path.name, page_count, units)


def inspect_document(document: Document) -> dict:
    """Bounded previews for every physical page, including empty ones."""
    by_page = {unit.page_or_slide: unit.text for unit in document.units}
    remaining = MAX_TEXT_CHARS
    pages = []
    truncated = False
    for page in range(1, document.page_count + 1):
        text = by_page.get(page, "")
        preview = text[:remaining]
        remaining -= len(preview)
        omitted = len(preview) < len(text)
        truncated = truncated or omitted
        pages.append({"page_or_slide": page, "text": preview, "truncated": omitted})
    result = {
        "type": "document",
        "filename": document.filename,
        "page_count": document.page_count,
        "pages": pages,
        "truncated": truncated,
        "synthetic": document.synthetic,
    }
    if document.synthetic:
        result["suggested_question"] = SAMPLE_QUESTION
    return result


def _selected_chunks(document: Document, pages: object) -> list[RetrievedChunk]:
    from gct.grounder.answer import _build_labeled_context
    from gct.ingest.chunk import chunk_units
    from gct.retriever.retrieve import RetrievedChunk

    if pages is _UNSET:
        if document.page_count > MAX_SELECTED_PAGES:
            raise ProofError("pages_required", "Select up to five pages or slides first.")
        selected = set(range(1, document.page_count + 1))
    else:
        if (
            not isinstance(pages, list)
            or not 1 <= len(pages) <= MAX_SELECTED_PAGES
            or any(type(page) is not int or not 1 <= page <= document.page_count for page in pages)
            or len(set(pages)) != len(pages)
        ):
            raise ProofError("invalid_pages", "Select one to five distinct valid page numbers.")
        selected = set(pages)
    units = [unit for unit in document.units if unit.page_or_slide in selected]
    if sum(len(unit.text) for unit in units) > MAX_TEXT_CHARS:
        raise ProofError("context_too_large", "Selected text exceeds 40,000 characters.")
    chunks = []
    page_indices: dict[int, int] = {}
    for chunk in chunk_units(units):
        index = page_indices.get(chunk.page_or_slide, 0)
        page_indices[chunk.page_or_slide] = index + 1
        identity = json.dumps(
            [chunk.file, chunk.page_or_slide, index, chunk.text], ensure_ascii=True
        ).encode("utf-8")
        chunks.append(
            RetrievedChunk(
                chunk_id=hashlib.sha256(identity).hexdigest(),
                text=chunk.text,
                file=chunk.file,
                page_or_slide=chunk.page_or_slide,
                score=1.0,
            )
        )
    # Include overlap and source-label overhead in the actual context bound.
    context, _ = _build_labeled_context(chunks)
    if len(context) > MAX_TEXT_CHARS:
        raise ProofError("context_too_large", "Selected context exceeds 40,000 characters.")
    return chunks


class StdioGeneration:
    """Delegate only text generation; all model output returns through Grounder."""

    model_id = "local-proof-peer"

    def __init__(self, source: TextIO, destination: TextIO) -> None:
        self.source = source
        self.destination = destination

    def generate(self, messages: Sequence[Message]) -> str:
        _emit(self.destination, {"type": "generate", "messages": list(messages)})
        try:
            response = _read_object(self.source)
        except ProofError as exc:
            raise GenerationFailure(
                "The generation connection returned an invalid response."
            ) from exc
        if response.get("type") == "generation_error":
            # Never interpolate peer-controlled error codes, messages, paths, or tokens.
            raise GenerationFailure("The generation service could not complete this request.")
        if (
            response.get("type") != "generated"
            or not isinstance(response.get("text"), str)
            or len(response["text"]) > MAX_TEXT_CHARS
        ):
            raise GenerationFailure("The generation connection returned an invalid response.")
        return response["text"]


def run(source: TextIO, destination: TextIO, *, inspect: bool = False) -> int:
    """Serve one request. Only the two supplied streams carry the bridge protocol."""
    # The existing Grounder imports config indirectly; config eagerly loads .env.
    # This peer does not use credentials, so prevent that incidental file read.
    previous_dotenv = os.environ.get("PYTHON_DOTENV_DISABLED")
    os.environ["PYTHON_DOTENV_DISABLED"] = "1"
    old_logging = logging.root.manager.disable
    logging.disable(sys.maxsize)
    try:
        with (
            contextlib.redirect_stdout(_DiscardOutput()),
            contextlib.redirect_stderr(_DiscardOutput()),
        ):
            request = _read_object(source)
            allowed = (
                {"sample", "file_path"} if inspect else {"sample", "file_path", "question", "pages"}
            )
            if set(request) - allowed:
                raise ProofError("invalid_request", "The request contains unsupported fields.")
            if not inspect:
                question = request.get("question")
                if (
                    not isinstance(question, str)
                    or not question.strip()
                    or len(question) > MAX_QUESTION_CHARS
                ):
                    raise ProofError(
                        "invalid_question", "Enter a question of 1 to 2,000 characters."
                    )
            document = _load_document(request)
            if inspect:
                _emit(destination, inspect_document(document))
                return 0
            from gct.grounder.answer import answer

            chunks = _selected_chunks(document, request.get("pages", _UNSET))
            result = answer(
                question,
                chunks,
                "local-proof",
                generator=StdioGeneration(source, destination),
            )
            _emit(destination, {"type": "result", "result": asdict(result)})
            return 0
    except ProofError as exc:
        _emit(destination, {"type": "error", "code": exc.code, "message": exc.message})
        return 1
    except Exception:
        _emit(
            destination,
            {"type": "error", "code": "internal", "message": "The local proof could not finish."},
        )
        return 1
    finally:
        logging.disable(old_logging)
        if previous_dotenv is None:
            os.environ.pop("PYTHON_DOTENV_DISABLED", None)
        else:
            os.environ["PYTHON_DOTENV_DISABLED"] = previous_dotenv


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--inspect", action="store_true", help="Return bounded document previews.")
    options = parser.parse_args()
    return run(sys.stdin, sys.stdout, inspect=options.inspect)


if __name__ == "__main__":
    raise SystemExit(main())
