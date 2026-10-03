"""Offline proof of the stdio boundary, source provenance, and bounded context."""

from __future__ import annotations

import io
import json
import os
import subprocess
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pypdf import PdfWriter
from reportlab.pdfgen.canvas import Canvas

from gct import local_proof
from gct.ingest.parse import ParsedUnit


def exchange(request, *replies, inspect=False):
    source = io.StringIO("\n".join(json.dumps(item) for item in (request, *replies)) + "\n")
    destination = io.StringIO()
    status = local_proof.run(source, destination, inspect=inspect)
    return status, [json.loads(line) for line in destination.getvalue().splitlines()]


def pdf(tmp_path, texts, name="Lecture notes.pdf"):
    path = tmp_path / name
    canvas = Canvas(str(path))
    for text in texts:
        if text:
            canvas.drawString(72, 720, text)
        canvas.showPage()
    canvas.save()
    return path


def test_sample_is_explicit_and_needs_no_selected_file():
    status, events = exchange({"sample": True}, inspect=True)
    assert status == 0
    document = events[0]
    assert document["synthetic"] is True
    assert document["filename"] == "local-proof-sample.pdf"
    assert document["page_count"] == 1
    assert document["suggested_question"] == local_proof.SAMPLE_QUESTION
    assert "Synthetic educational sample" in document["pages"][0]["text"]


def test_generation_reuses_provenance_and_validates_stale_labels(tmp_path):
    path = pdf(tmp_path, ["First page.", "", "Spaced practice spreads study over time."])
    request = {"file_path": str(path), "pages": [3], "question": "What is spaced practice?"}
    status, events = exchange(
        request,
        {
            "type": "generated",
            "text": "Spaced practice spreads sessions out [S2].\nCOVERAGE: complete",
        },
        {
            "type": "generated",
            "text": "Spaced practice spreads sessions out [S1].\nCOVERAGE: complete",
        },
    )
    assert status == 0
    assert [event["type"] for event in events] == ["generate", "generate", "result"]
    assert events[0]["messages"] == events[1]["messages"]
    context = events[0]["messages"][1]["content"]
    assert "[S1] (Lecture notes.pdf, p.3)" in context
    assert "First page." not in context
    result = events[-1]["result"]
    assert result["state"] == "GROUNDED"
    citation = result["citations"][0]
    assert citation["file"] == path.name
    assert citation["page_or_slide"] == 3
    assert len(citation["chunk_id"]) == 64
    _, repeated = exchange(
        request, {"type": "generated", "text": "Spaced sessions [S1].\nCOVERAGE: complete"}
    )
    assert repeated[-1]["result"]["citations"][0]["chunk_id"] == citation["chunk_id"]
    _, expanded = exchange(
        {**request, "pages": [1, 3]},
        {"type": "generated", "text": "Spaced sessions [S2].\nCOVERAGE: complete"},
    )
    assert expanded[-1]["result"]["citations"][0]["chunk_id"] == citation["chunk_id"]


def test_invalid_labels_stop_after_existing_two_attempt_budget():
    invalid = {"type": "generated", "text": "Unsupported label [S999].\nCOVERAGE: complete"}
    _, events = exchange({"sample": True, "question": "What is active recall?"}, invalid, invalid)
    assert [event["type"] for event in events] == ["generate", "generate", "result"]
    assert events[-1]["result"]["state"] == "INTEGRITY_FLAGGED"
    assert events[-1]["result"]["citations"] == []


def test_refusal_is_not_a_transport_error():
    _, events = exchange(
        {"sample": True, "question": "What is the exact optimal interval?"},
        {"type": "generated", "text": "COVERAGE: gaps: the exact optimal interval"},
    )
    result = events[-1]["result"]
    assert result["state"] == "REFUSAL"
    assert result["error"] is None
    assert result["citations"] == []


@pytest.mark.parametrize(
    "reply",
    [
        {"type": "generation_error", "code": "SECRET raw provider message /private/path"},
        {"type": "other", "text": "SECRET"},
        {"type": "generated", "text": None},
        {"type": "generated", "text": "SECRET" * 10_000},
    ],
)
def test_sdk_and_protocol_errors_are_terminal_and_redacted(reply):
    status, events = exchange({"sample": True, "question": "What is active recall?"}, reply)
    assert status == 0
    assert [event["type"] for event in events] == ["generate", "result"]
    result = events[-1]["result"]
    assert result["state"] == "ERROR"
    assert result["error"]["kind"] == "provider_terminal"
    assert "SECRET" not in json.dumps(result)
    assert "/private/path" not in json.dumps(result)


def test_peer_eof_is_terminal_without_retry():
    _, events = exchange({"sample": True, "question": "What is active recall?"})
    assert [event["type"] for event in events] == ["generate", "result"]
    assert events[-1]["result"]["state"] == "ERROR"


def test_inspection_preserves_empty_and_trailing_pages(tmp_path):
    path = pdf(tmp_path, ["One.", "", "Three.", ""])
    status, events = exchange({"file_path": str(path)}, inspect=True)
    assert status == 0
    document = events[0]
    assert document["page_count"] == 4
    assert [page["page_or_slide"] for page in document["pages"]] == [1, 2, 3, 4]
    assert document["pages"][1]["text"] == document["pages"][3]["text"] == ""
    assert document["truncated"] is False
    _, events = exchange({"file_path": str(path), "pages": [2], "question": "Anything here?"})
    assert [event["type"] for event in events] == ["result"]
    assert events[0]["result"]["state"] == "REFUSAL"


def test_pptx_provenance_and_notes_use_existing_parser(tmp_path):
    path = tmp_path / "Practice slides.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Active recall"
    slide.placeholders[1].text = "Retrieve from memory."
    slide.notes_slide.notes_text_frame.text = "Review with spaced practice."
    deck.save(str(path))
    _, inspected = exchange({"file_path": str(path)}, inspect=True)
    assert "Review with spaced practice." in inspected[0]["pages"][0]["text"]
    _, events = exchange(
        {"file_path": str(path), "question": "What is active recall?"},
        {"type": "generated", "text": "Retrieve from memory [S1].\nCOVERAGE: complete"},
    )
    assert events[-1]["result"]["citations"][0]["file"] == path.name
    assert events[-1]["result"]["citations"][0]["page_or_slide"] == 1


@pytest.mark.parametrize(
    "pages", [None, [], [0], [7], [True], [1, 1], [1, 2, 3, 4, 5, 6], "1", [{}]]
)
def test_invalid_page_selections_never_generate(tmp_path, pages):
    path = pdf(tmp_path, ["Page text."] * 6)
    status, events = exchange({"file_path": str(path), "pages": pages, "question": "Question?"})
    assert status == 1
    assert len(events) == 1
    assert events[0]["code"] == "invalid_pages"


def test_long_document_requires_explicit_pages_before_generation(tmp_path):
    path = pdf(tmp_path, ["Page text."] * 6)
    status, events = exchange({"file_path": str(path), "question": "Question?"})
    assert status == 1
    assert events[0]["code"] == "pages_required"


@pytest.mark.parametrize("question", [None, "", " \n ", 42, "q" * 2001])
def test_question_validation_precedes_reading_file(tmp_path, question):
    status, events = exchange({"file_path": str(tmp_path / "absent.pdf"), "question": question})
    assert status == 1
    assert events[0]["code"] == "invalid_question"


def test_file_size_and_page_count_limits_precede_full_extraction(tmp_path):
    large = tmp_path / "too-big.pdf"
    with large.open("wb") as stream:
        stream.truncate(local_proof.MAX_FILE_BYTES + 1)
    _, events = exchange({"file_path": str(large)}, inspect=True)
    assert events[0]["code"] == "file_too_large"
    many = tmp_path / "too-many.pdf"
    writer = PdfWriter()
    for _ in range(local_proof.MAX_PAGES + 1):
        writer.add_blank_page(width=10, height=10)
    writer.write(many)
    _, events = exchange({"file_path": str(many)}, inspect=True)
    assert events[0]["code"] == "too_many_pages"


def test_inspection_previews_are_capped_with_explicit_truncation():
    document = local_proof.Document(
        "notes.pdf",
        3,
        [ParsedUnit("a" * 30_000, "notes.pdf", 1), ParsedUnit("b" * 30_000, "notes.pdf", 3)],
    )
    inspected = local_proof.inspect_document(document)
    assert sum(len(page["text"]) for page in inspected["pages"]) == 40_000
    assert inspected["truncated"] is True
    assert inspected["pages"][1]["text"] == ""
    assert inspected["pages"][2]["truncated"] is True


@pytest.mark.parametrize("text", ["x" * 40_001, "word " * 7_000])
def test_text_and_actual_overlapped_context_limits(monkeypatch, text):
    document = local_proof.Document("notes.pdf", 1, [ParsedUnit(text, "notes.pdf", 1)])
    monkeypatch.setattr(local_proof, "_load_document", lambda _: document)
    status, events = exchange({"sample": True, "question": "Question?"})
    assert status == 1
    assert len(events) == 1
    assert events[0]["code"] == "context_too_large"


def test_corrupt_file_and_parser_diagnostics_never_leak(tmp_path, capsys):
    path = tmp_path / "secret-name.pdf"
    path.write_bytes(b"SECRET malformed document")
    _, events = exchange({"file_path": str(path)}, inspect=True)
    assert events[0]["code"] == "unparseable"
    assert "SECRET" not in json.dumps(events)
    assert "secret-name" not in json.dumps(events)
    assert capsys.readouterr() == ("", "")


@pytest.mark.parametrize("raw", ["{oops}\n", "[]\n", "", "x" * (local_proof.MAX_LINE_CHARS + 1)])
def test_invalid_first_protocol_line(raw):
    destination = io.StringIO()
    assert local_proof.run(io.StringIO(raw), destination) == 1
    assert json.loads(destination.getvalue())["code"] == "invalid_protocol"


def test_cli_does_not_read_dotenv_connect_or_emit_diagnostics(tmp_path):
    # A separate process proves the import path, unlike this test process whose
    # shared conftest has already imported the database configuration.
    script = """
import sys
def audit(event, args):
    if event == 'open' and str(args[0]).endswith('.env'):
        raise RuntimeError('UNEXPECTED_CREDENTIAL_READ')
    if event == 'socket.connect':
        raise RuntimeError('UNEXPECTED_NETWORK')
sys.addaudithook(audit)
from gct.local_proof import main
raise SystemExit(main())
"""
    root = Path(__file__).resolve().parents[2]
    environment = {**os.environ, "PYTHONPATH": str(root / "src")}
    completed = subprocess.run(
        [sys.executable, "-c", script],
        input=(
            json.dumps({"sample": True, "question": "What is active recall?"})
            + "\n"
            + json.dumps(
                {"type": "generated", "text": "Retrieve from memory [S1].\nCOVERAGE: complete"}
            )
            + "\n"
        ),
        text=True,
        capture_output=True,
        cwd=tmp_path,
        env=environment,
        timeout=10,
        check=False,
    )
    assert completed.returncode == 0, completed.stderr
    assert completed.stderr == ""
    events = [json.loads(line) for line in completed.stdout.splitlines()]
    assert [event["type"] for event in events] == ["generate", "result"]
    assert events[-1]["result"]["state"] == "GROUNDED"
